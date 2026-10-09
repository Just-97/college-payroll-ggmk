import os
import sys
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import docx

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
LABS_DIR = os.path.join(BASE_DIR, 'лабы')
SRC_DIR = os.path.join(BASE_DIR, 'src')
ASSETS_DIR = os.path.join(LABS_DIR, 'assets')
sys.path.append(os.path.dirname(__file__))
from generate_labs_helpers import add_header, add_code_block, add_paragraph_text, add_heading_2


def read_file(rel_path):
    with open(os.path.join(SRC_DIR, rel_path), 'r', encoding='utf-8') as f:
        return f.read()


# ==============================================================================
# 1. ЛАБОРАТОРНАЯ РАБОТА № 27: MVC
# ==============================================================================
print("Generating Lab 27 (MVC)...")
doc27 = docx.Document()

# Margins
sections = doc27.sections
for s in sections:
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(0.75)

add_header(
    doc27,
    lab_num="27",
    theme="Исследование программной системы, на основе использования шаблона MVC",
    goal="Научиться разрабатывать программную систему с использованием MVC шаблона в рамках дисциплины ТРПО."
)

add_heading_2(doc27, "1. Структурная схема (UML-диаграмма классов MVC)")
add_paragraph_text(
    doc27,
    "В соответствии с заданием было выполнено проектирование модуля администрирования пользователей веб-приложения «Расчет заработной платы преподавателей УО «Гомельский государственный машиностроительный колледж» на основе архитектурного шаблона Model-View-Controller (MVC). На рисунке 1 представлена UML-диаграмма классов спроектированной подсистемы."
)

p_img1 = doc27.add_paragraph()
p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img1.paragraph_format.space_before = Pt(6)
p_img1.paragraph_format.space_after = Pt(2)
doc27.add_picture(os.path.join(ASSETS_DIR, 'uml_mvc.png'), width=Inches(6.2))

p_cap1 = doc27.add_paragraph()
p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cap1.paragraph_format.space_after = Pt(8)
r_cap1 = p_cap1.add_run("Рисунок 1 — UML-диаграмма классов модуля администрирования (паттерн MVC)")
r_cap1.font.name = 'Times New Roman'
r_cap1.font.size = Pt(11)
r_cap1.italic = True

add_heading_2(doc27, "2. Описание компонентов спроектированной архитектуры MVC")
add_paragraph_text(
    doc27,
    "1. Model (UserModel.ts) — модель данных предметной области. Является единственным источником истины о пользователях системы. Реализует паттерн «Наблюдатель» (Observer) через методы subscribe() и notify(), инкапсулирует бизнес-валидацию входных данных (длина логина, сложность пароля, уникальность email, корректность роли) и осуществляет асинхронную синхронизацию с облачной СУБД Supabase (PostgreSQL) и локальным хранилищем.",
    bold_prefix="• "
)
add_paragraph_text(
    doc27,
    "2. Controller (AdminController.ts) — контроллер бизнес-операций. Выступает связующим звеном между представлением и моделью. Перехватывает намерения пользователя (создание, редактирование, удаление учетной записи, блокировка доступа, генерация стойких паролей), валидирует параметры запроса и командует Модели применить изменения.",
    bold_prefix="• "
)
add_paragraph_text(
    doc27,
    "3. View (AdminView.tsx) — пользовательский графический интерфейс администратора колледжа. Подписывается на Модель (hook useEffect) и пассивно отображает текущий список пользователей, фильтры и статус БД. При действиях пользователя (клик по кнопкам, отправка форм) делегирует управление Контроллеру, исключая прямую модификацию состояния внутри UI.",
    bold_prefix="• "
)

add_heading_2(doc27, "3. Исходный код программы")
add_code_block(doc27, "Листинг 1 — Модель данных пользователей (UserModel.ts)", read_file('patterns/mvc/UserModel.ts'))
add_code_block(doc27, "Листинг 2 — Контроллер администратора (AdminController.ts)", read_file('patterns/mvc/AdminController.ts'))

# Сокращенная презентационная часть View для читаемости
view_code = read_file('patterns/mvc/AdminView.tsx')
if len(view_code) > 4000:
    view_code_short = view_code[:3800] + "\n\n// ... [Продолжение разметки модальных окон и таблицы пользователей] ...\n"
else:
    view_code_short = view_code
add_code_block(doc27, "Листинг 3 — Представление панели администратора (AdminView.tsx, фрагмент)", view_code_short)

add_heading_2(doc27, "4. Результат работы программы")
add_paragraph_text(
    doc27,
    "Разработанный модуль был интегрирован в веб-приложение колледжа на отдельном изолированном маршруте (/#/admin). На рисунке 2 представлен внешний вид панели администратора: отображение таблицы учетных записей (администратор, бухгалтер, преподаватели), индикация ролей и паролей, фильтрация по ФИО/логину, а также статус подключения к облачной СУБД Supabase (PostgreSQL)."
)

p_img2 = doc27.add_paragraph()
p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img2.paragraph_format.space_before = Pt(6)
p_img2.paragraph_format.space_after = Pt(2)
doc27.add_picture(os.path.join(ASSETS_DIR, 'res_mvc.png'), width=Inches(6.2))

p_cap2 = doc27.add_paragraph()
p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cap2.paragraph_format.space_after = Pt(8)
r_cap2 = p_cap2.add_run("Рисунок 2 — Результат работы модуля администрирования на основе паттерна MVC")
r_cap2.font.name = 'Times New Roman'
r_cap2.font.size = Pt(11)
r_cap2.italic = True

add_heading_2(doc27, "Вывод")
add_paragraph_text(
    doc27,
    "В ходе выполнения лабораторной работы были исследованы принципы проектирования программных систем на основе шаблона Model-View-Controller (MVC). С использованием языка UML разработана структурная схема классов модуля администрирования. Реализовано четкое разделение ответственности: бизнес-логика валидации и хранения пользователей инкапсулирована в UserModel, обработка команд вынесена в AdminController, а визуализация закреплена за AdminView. Внедрение паттерна Observer позволило обеспечить слабую связность компонентов и реактивное обновление интерфейса при модификации данных."
)

target_path27 = os.path.join(LABS_DIR, 'Лабораторная работа № 27.docx')
doc27.save(target_path27)
print(f"Saved: {target_path27}")

# ==============================================================================
# 2. ЛАБОРАТОРНАЯ РАБОТА № 28: MVP
# ==============================================================================
print("Generating Lab 28 (MVP)...")
doc28 = docx.Document()
for s in doc28.sections:
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(0.75)

add_header(
    doc28,
    lab_num="28",
    theme="Исследование программной системы, на основе использования шаблона MVP",
    goal="Научиться разрабатывать программную среду с использованием MVP шаблона в рамках дисциплины ТРПО."
)

add_heading_2(doc28, "1. Структурная схема (UML-диаграмма классов MVP)")
add_paragraph_text(
    doc28,
    "В соответствии с заданием спроектирован модуль персонального расчетного листка преподавателя колледжа (согласно требованиям ст. 80 Трудового кодекса Республики Беларусь) с применением архитектурного шаблона Model-View-Presenter (MVP, подтип Passive View). На рисунке 1 представлена UML-диаграмма классов с выделением интерфейсного контракта представления."
)

p_img28 = doc28.add_paragraph()
p_img28.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img28.paragraph_format.space_before = Pt(6)
p_img28.paragraph_format.space_after = Pt(2)
doc28.add_picture(os.path.join(ASSETS_DIR, 'uml_mvp.png'), width=Inches(6.2))

p_cap28 = doc28.add_paragraph()
p_cap28.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cap28.paragraph_format.space_after = Pt(8)
r_cap28 = p_cap28.add_run("Рисунок 1 — UML-диаграмма классов расчетного листка преподавателя (паттерн MVP)")
r_cap28.font.name = 'Times New Roman'
r_cap28.font.size = Pt(11)
r_cap28.italic = True

add_heading_2(doc28, "2. Описание компонентов спроектированной архитектуры MVP")
add_paragraph_text(
    doc28,
    "1. Contract Interface (IPayslipView.ts) — строгий интерфейс представления (View Interface). Задает контракт: методы setPeriodOptions(), renderPayslip(dto), showNotification() и executeBrowserPrint(). Презентер зависит исключительно от этого интерфейса и не имеет прямых связей с библиотекой React или DOM-деревом браузера, что обеспечивает 100% покрытие модульными тестами.",
    bold_prefix="• "
)
add_paragraph_text(
    doc28,
    "2. Presenter (PayslipPresenter.ts) — представитель и единственный координатор логики выплат. Извлекает сырые финансовые данные из модели, выполняет расчеты по ставкам законодательства РБ (подоходный налог 13%, ФСЗН 1%, профсоюз 1%, надбавки за категорию, стаж и кураторство), форматирует денежные единицы (BYN) и командует пассивному виду через контрактный интерфейс.",
    bold_prefix="• "
)
add_paragraph_text(
    doc28,
    "3. Model (PayrollModel.ts) — модель доменных данных преподавателя. Хранит информацию о квалификационной категории, педагогическом стаже, заведовании кабинетами, часовой нагрузке и истории начислений по месяцам.",
    bold_prefix="• "
)
add_paragraph_text(
    doc28,
    "4. View (PayslipView.tsx) — пассивное представление (Passive View). Реализует методы IPayslipView. Не содержит логики вычисления налогов или форматирования — при кликах на выбор месяца или печать А4 моментально делегирует событие методам Презентера.",
    bold_prefix="• "
)

add_heading_2(doc28, "3. Исходный код программы")
add_code_block(doc28, "Листинг 1 — Контракт представления (IPayslipView.ts)", read_file('patterns/mvp/IPayslipView.ts'))
add_code_block(doc28, "Листинг 2 — Доменная модель расчетного листка (PayrollModel.ts)", read_file('patterns/mvp/PayrollModel.ts'))
add_code_block(doc28, "Листинг 3 — Презентер расчетного листка (PayslipPresenter.ts)", read_file('patterns/mvp/PayslipPresenter.ts'))

pv_code = read_file('patterns/mvp/PayslipView.tsx')
if len(pv_code) > 3500:
    pv_code_short = pv_code[:3300] + "\n\n// ... [Продолжение JSX-верстки пассивного представления] ...\n"
else:
    pv_code_short = pv_code
add_code_block(doc28, "Листинг 4 — Пассивное представление расчетного листка (PayslipView.tsx, фрагмент)", pv_code_short)

add_heading_2(doc28, "4. Результат работы программы")
add_paragraph_text(
    doc28,
    "На рисунке 2 представлен результат работы спроектированного модуля MVP. Преподаватель видит официальный расчетный листок за выбранный месяц (Сентябрь 2026) с детальным разделением: начисления (оклад 76 ч, специфика 40%, надбавка за высшую категорию, стаж 16 лет, кураторство группы 24-ПО-1, заведование кабинетом), удержания (подоходный налог 13%, взносы ФСЗН 1%, профсоюзный взнос 1%, аванс) и итоговая сумма к выплате на карт-счет. По нажатию кнопки вызывается метод презентера для печати листка под формат А4."
)

p_img28_res = doc28.add_paragraph()
p_img28_res.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img28_res.paragraph_format.space_before = Pt(6)
p_img28_res.paragraph_format.space_after = Pt(2)
doc28.add_picture(os.path.join(ASSETS_DIR, 'res_mvp.png'), width=Inches(6.2))

p_cap28_res = doc28.add_paragraph()
p_cap28_res.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cap28_res.paragraph_format.space_after = Pt(8)
r_cap28_res = p_cap28_res.add_run("Рисунок 2 — Результат работы модуля расчетного листка преподавателя (паттерн MVP)")
r_cap28_res.font.name = 'Times New Roman'
r_cap28_res.font.size = Pt(11)
r_cap28_res.italic = True

add_heading_2(doc28, "Вывод")
add_paragraph_text(
    doc28,
    "В ходе лабораторной работы освоен паттерн Model-View-Presenter (MVP) в конфигурации Passive View. Благодаря выделению интерфейса IPayslipView достигнута полная независимость бизнес-логики формирования расчетного листка от платформы отображения. Презентер инкапсулирует сложные финансовые расчеты в строгом соответствии со ст. 80 ТК РБ, что упрощает масштабирование системы, сопровождение и проведение модульного автоматизированного тестирования без привлечения браузерного окружения."
)

target_path28 = os.path.join(LABS_DIR, 'Лабораторная работа № 28.docx')
doc28.save(target_path28)
print(f"Saved: {target_path28}")

# ==============================================================================
# 3. ЛАБОРАТОРНАЯ РАБОТА № 29: MVVM
# ==============================================================================
print("Generating Lab 29 (MVVM)...")
doc29 = docx.Document()
for s in doc29.sections:
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(0.75)

add_header(
    doc29,
    lab_num="29",
    theme="Исследование программной системы, на основе использования шаблона MVVM",
    goal="Научиться разрабатывать программную среду с использованием MVVM шаблона в рамках дисциплины ТРПО."
)

add_heading_2(doc29, "1. Структурная схема (UML-диаграмма классов MVVM)")
add_paragraph_text(
    doc29,
    "В рамках работы спроектировано автоматизированное рабочее место (АРМ) главного бухгалтера колледжа для ведения сводной ведомости, расчета стимулирующих надплат и экспорта финансовой отчетности на основе паттерна Model-View-ViewModel (MVVM). На рисунке 1 представлена UML-диаграмма классов подсистемы."
)

p_img29 = doc29.add_paragraph()
p_img29.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img29.paragraph_format.space_before = Pt(6)
p_img29.paragraph_format.space_after = Pt(2)
doc29.add_picture(os.path.join(ASSETS_DIR, 'uml_mvvm.png'), width=Inches(6.2))

p_cap29 = doc29.add_paragraph()
p_cap29.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cap29.paragraph_format.space_after = Pt(8)
r_cap29 = p_cap29.add_run("Рисунок 1 — UML-диаграмма классов рабочего места бухгалтера (паттерн MVVM)")
r_cap29.font.name = 'Times New Roman'
r_cap29.font.size = Pt(11)
r_cap29.italic = True

add_heading_2(doc29, "2. Описание компонентов спроектированной архитектуры MVVM")
add_paragraph_text(
    doc29,
    "1. Model (AccountantModel.ts) — доменная модель начислений и преподавателей. Содержит асинхронные методы выборки и модификации данных (loadAll(), savePayrollRecord(), deletePayrollRecord(), saveTeacher(), deleteTeacher()) с автоматической синхронизацией в облачную PostgreSQL (Supabase).",
    bold_prefix="• "
)
add_paragraph_text(
    doc29,
    "2. ViewModel (AccountantViewModel.ts) — реактивная модель представления (Custom React Hook). Предоставляет свойства состояния для двустороннего связывания (Two-Way Data Binding: searchQuery, departmentFilter, selectedPeriod), вычисляемые аналитические свойства через useMemo (filteredRecords, financialSummary — автоматический пересчет общего фонда зарплаты колледжа, налогов и выплат) и инкапсулированные команды (Commands: savePayrollCommand, deletePayrollCommand, saveTeacherCommand, exportCsvCommand).",
    bold_prefix="• "
)
add_paragraph_text(
    doc29,
    "3. View (AccountantView.tsx) — декларативное представление рабочего места бухгалтера. Связывается со свойствами ViewModel декларативно через синтаксис JSX, отображает аналитические карточки, фильтры и таблицу начислений, и триггерит выполнение команд при взаимодействии с пользователем.",
    bold_prefix="• "
)

add_heading_2(doc29, "3. Исходный код программы")
add_code_block(doc29, "Листинг 1 — Доменная модель бухгалтера (AccountantModel.ts)", read_file('patterns/mvvm/AccountantModel.ts'))
add_code_block(doc29, "Листинг 2 — Модель представления (AccountantViewModel.ts)", read_file('patterns/mvvm/AccountantViewModel.ts'))

av_code = read_file('patterns/mvvm/AccountantView.tsx')
if len(av_code) > 3500:
    av_code_short = av_code[:3300] + "\n\n// ... [Продолжение декларативного связывания View и модальных окон] ...\n"
else:
    av_code_short = av_code
add_code_block(doc29, "Листинг 3 — Декларативное представление бухгалтера (AccountantView.tsx, фрагмент)", av_code_short)

add_heading_2(doc29, "4. Результат работы программы")
add_paragraph_text(
    doc29,
    "На рисунке 2 продемонстрирован интерфейс АРМ бухгалтера УО «ГГМК». В верхней части отображаются вычисляемые свойства ViewModel: суммарный фонд оплаты труда (ФОТ), удержания подоходного налога и ФСЗН, итого к выплате по колледжу. Ниже расположены элементы управления, привязанные двусторонним связыванием (поиск преподавателя, фильтр по цикловым комиссиям), и интерактивная таблица начислений с возможностью экспорта ведомости в Excel/CSV по нажатию на кнопку команды."
)

p_img29_res = doc29.add_paragraph()
p_img29_res.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img29_res.paragraph_format.space_before = Pt(6)
p_img29_res.paragraph_format.space_after = Pt(2)
doc29.add_picture(os.path.join(ASSETS_DIR, 'res_mvvm.png'), width=Inches(6.2))

p_cap29_res = doc29.add_paragraph()
p_cap29_res.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cap29_res.paragraph_format.space_after = Pt(8)
r_cap29_res = p_cap29_res.add_run("Рисунок 2 — Результат работы АРМ бухгалтера на основе паттерна MVVM")
r_cap29_res.font.name = 'Times New Roman'
r_cap29_res.font.size = Pt(11)
r_cap29_res.italic = True

add_heading_2(doc29, "Вывод")
add_paragraph_text(
    doc29,
    "В ходе выполнения лабораторной работы исследована архитектурная концепция Model-View-ViewModel (MVVM). Использование модели представления позволило полностью отделить визуальную часть от механизмов финансовой фильтрации и расчета показателей колледжа. Механизм Data Binding освобождает код View от ручного обновления DOM-элементов при изменениях поисковой строки или списков отделений, а инкапсуляция логики сохранения и экспорта в команды (Commands) обеспечивает высокую надежность, переиспользуемость и тестируемость программной системы."
)

target_path29 = os.path.join(LABS_DIR, 'Лабораторная работа № 29.docx')
doc29.save(target_path29)
print(f"Saved: {target_path29}")

print("ALL LABS SUCCESSFULLY GENERATED!")
