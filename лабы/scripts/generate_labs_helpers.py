import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header(doc, lab_num, theme, goal):
    p = doc.add_paragraph()
    r = p.add_run(f"Лабораторная работа №{lab_num}\n")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    
    info_items = [
        ("Номер учебной группы: ", "ПО-41"),
        ("Фамилия, инициалы обучающегося: ", "____________________"),
        ("Дата выполнения работы: ", "10.10.2026"),
        ("Тема работы: ", theme),
        ("Цель работы: ", goal),
        ("Задание:", "")
    ]
    
    for title, val in info_items:
        p_info = doc.add_paragraph()
        p_info.paragraph_format.space_after = Pt(2)
        p_info.paragraph_format.space_before = Pt(0)
        p_info.paragraph_format.line_spacing = 1.15
        r_t = p_info.add_run(title)
        r_t.bold = True
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(14)
        if val:
            r_v = p_info.add_run(val)
            r_v.font.name = 'Times New Roman'
            r_v.font.size = Pt(14)
            
    tasks = [
        f"1. Схематично описать {theme.split('шаблона ')[-1]} паттерн с помощью UML языка",
        f"2. Выполнить проектирование программы с использованием {theme.split('шаблона ')[-1]} паттерна"
    ]
    for t in tasks:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.left_indent = Inches(0.25)
        p_t.paragraph_format.space_after = Pt(2)
        p_t.paragraph_format.line_spacing = 1.15
        r_task = p_t.add_run(t)
        r_task.font.name = 'Times New Roman'
        r_task.font.size = Pt(14)

    p_res = doc.add_paragraph()
    p_res.paragraph_format.space_before = Pt(8)
    p_res.paragraph_format.space_after = Pt(4)
    r_res = p_res.add_run("Результат выполнения работы:")
    r_res.bold = True
    r_res.font.name = 'Times New Roman'
    r_res.font.size = Pt(14)

def add_code_block(doc, title, code_text):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(2)
    r_t = p_title.add_run(title)
    r_t.bold = True
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(12)
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, 'F1F5F9')
    set_cell_margins(cell, 120, 120, 150, 150)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    lines = code_text.strip().split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line)
        run.font.name = 'Courier New'
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(15, 23, 42)

def add_paragraph_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(14)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)

print("Helper functions defined successfully")
